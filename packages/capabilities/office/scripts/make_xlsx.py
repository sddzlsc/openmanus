#!/usr/bin/env python3
"""Generate a spreadsheet deliverable (and its CSV companion). Requires openpyxl."""

from __future__ import annotations

import csv
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _common import ensure_out, fail, parse_args  # noqa: E402

ROWS = [
    ["阶段", "工作项", "负责人", "开始日期", "结束日期", "状态", "完成度"],
    ["1", "需求澄清", "Agent", "D1", "D2", "已完成", "100%"],
    ["2", "资料收集", "Agent", "D2", "D3", "已完成", "100%"],
    ["3", "初稿生成", "Agent", "D3", "D5", "进行中", "60%"],
    ["4", "评审与修改", "用户", "D5", "D7", "待开始", "0%"],
    ["5", "最终交付", "Agent", "D7", "D8", "待开始", "0%"],
]


def main() -> None:
    args = parse_args("生成 Excel 表格")
    out = ensure_out(args.out)

    csv_path = os.path.join(out, "plan.csv")
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as handle:
        csv.writer(handle).writerows(ROWS)
    print(f"saved {csv_path}")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        fail("openpyxl 未安装（沙箱镜像应包含 openpyxl）")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "执行计划"
    for row in ROWS:
        sheet.append(row)
    header_fill = PatternFill("solid", fgColor="4338CA")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    widths = [8, 24, 12, 12, 12, 12, 10]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    summary = workbook.create_sheet("说明")
    summary["A1"] = args.title
    summary["A2"] = args.prompt
    summary.column_dimensions["A"].width = 80

    target = os.path.join(out, "workbook.xlsx")
    workbook.save(target)
    print(f"saved {target}")


if __name__ == "__main__":
    main()
