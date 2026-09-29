#!/usr/bin/env python3
"""Generate a Word document deliverable. Requires python-docx."""

from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _common import CJK_FONT_CANDIDATES, ensure_out, fail, parse_args, sections  # noqa: E402


def main() -> None:
    args = parse_args("生成 Word 文档")
    out = ensure_out(args.out)
    try:
        from docx import Document
        from docx.shared import Pt
        from docx.oxml.ns import qn
    except ImportError:
        fail("python-docx 未安装（沙箱镜像应包含 python-docx）")

    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(11)
    style.font.name = CJK_FONT_CANDIDATES[0]
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT_CANDIDATES[0])

    doc.add_heading(args.title, level=0)
    doc.add_paragraph("由 Wiwana 智能体生成 · 交付物类型：文档")
    doc.add_paragraph(args.prompt)

    for heading, paragraphs in sections(args.prompt):
        doc.add_heading(heading, level=1)
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)

    doc.add_heading("待确认事项", level=1)
    doc.add_paragraph("1. 目标受众与使用场景\n2. 交付格式与截止时间\n3. 是否需要品牌视觉规范")

    target = os.path.join(out, "deliverable.docx")
    doc.save(target)
    print(f"saved {target}")


if __name__ == "__main__":
    main()
